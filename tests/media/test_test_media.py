"""Sintetik to'plam (TZ 13.3) haqiqatan kutilgan xususiyatlarga ega: keyingi bosqich testlari shunga tayanadi."""

import json
from fractions import Fraction
from typing import Any

import pytest

from darsstudio.media import ffmpeg


def _streams(path: str) -> dict[str, dict[str, Any]]:
    r = ffmpeg.run(["-v", "error", "-count_packets", "-show_streams", "-of", "json", path], tool="ffprobe")
    return {s["codec_type"]: s for s in json.loads(r.stdout)["streams"]}


@pytest.mark.parametrize("name", ["green_key.mp4", "av_offset.mov", "ntsc_2997.mp4", "vfr.mp4", "timer.mp4"])
def test_video_fps_and_frames(test_media: dict[str, Any], name: str) -> None:
    exp = test_media[name]
    v = _streams(exp["path"])["video"]
    assert (v["width"], v["height"]) == (1920, 1080)
    assert Fraction(v["r_frame_rate"]) == Fraction(exp["fps"])
    assert int(v["nb_read_packets"]) == exp["nb_frames"]


def test_av_offset(test_media: dict[str, Any]) -> None:
    exp = test_media["av_offset.mov"]
    s = _streams(exp["path"])
    diff_ms = (float(s["video"]["start_time"]) - float(s["audio"]["start_time"])) * 1000
    assert diff_ms == pytest.approx(exp["video_minus_audio_start_ms"], abs=0.5)
    assert s["video"]["pix_fmt"] == "yuvj420p" and s["video"]["color_range"] == "pc"
    assert s["audio"]["codec_name"] == "pcm_s16le"


def test_ntsc_has_aac(test_media: dict[str, Any]) -> None:
    assert _streams(test_media["ntsc_2997.mp4"]["path"])["audio"]["codec_name"] == "aac"


def test_vfr_is_vfr(test_media: dict[str, Any]) -> None:
    v = _streams(test_media["vfr.mp4"]["path"])["video"]
    assert Fraction(v["avg_frame_rate"]) < Fraction(v["r_frame_rate"])


def test_beep_clip_has_audio(test_media: dict[str, Any]) -> None:
    a = _streams(test_media["beep.mov"]["path"])["audio"]
    assert int(a["sample_rate"]) == 48000


def test_pptx_content(test_media: dict[str, Any]) -> None:
    from pptx import Presentation

    exp = test_media["lesson.pptx"]["slides"]
    prs = Presentation(test_media["lesson.pptx"]["path"])
    assert len(prs.slides) == len(exp) == 12
    for s, e in zip(prs.slides, exp, strict=True):
        assert s.shapes.title.text == e["title"]
        assert s.notes_slide.notes_text_frame.text == e["notes"]
        has_anim = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}timing") is not None
        assert has_anim == e["animated"]


def test_run_raises_readable_error(tmp_path: Any) -> None:
    with pytest.raises(ffmpeg.FfmpegError, match="xato bilan tugadi"):
        ffmpeg.run(["-v", "error", "-i", str(tmp_path / "yoq.mp4")], tool="ffprobe")
