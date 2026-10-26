"""Sintetik test media (TZ 13.3). Repo'ga qo'shilmaydi, test boshida yaratiladi.

    python tools/make_test_media.py [OUT_DIR]

OUT_DIR/manifest.json da har fayl va testlar tekshiradigan kutilgan qiymatlar yoziladi.
"""

import json
import sys
from collections.abc import Callable
from pathlib import Path
from typing import Any

from darsstudio.media import ffmpeg

# Generator o'zgarsa oshiriladi: conftest eski to'plamni qayta yaratadi
VERSION = 2
DEFAULT_OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "generated"

HD = "1920x1080"
X264 = ["-c:v", "libx264", "-preset", "veryfast", "-crf", "18", "-pix_fmt", "yuv420p"]
BT709 = ["-colorspace", "bt709", "-color_primaries", "bt709", "-color_trc", "bt709"]
GREEN = "0x2EB84B"


def _ff(*args: str) -> None:
    ffmpeg.run(["-v", "error", "-nostdin", "-y", *args])


def green_key(out: Path) -> dict[str, Any]:
    """1. Yashil fon + harakatlanuvchi teksturali ellips ("odam") + o'ng chetda oq "devor" polosasi."""
    dur, fps = 4, 25
    pw, ph = 400, 640
    ellipse = f"if(lte(pow((X-{pw / 2})/{pw / 2}\\,2)+pow((Y-{ph / 2})/{ph / 2}\\,2)\\,1)\\,255\\,0)"
    graph = (
        f"color=c={GREEN}:s={HD}:r={fps}:d={dur}[bg];"
        f"color=c=0xC89878:s={pw}x{ph}:r={fps}:d={dur},noise=alls=25:allf=t,format=yuva420p,"
        f"geq=lum=lum(X\\,Y):cb=cb(X\\,Y):cr=cr(X\\,Y):a={ellipse}[p];"
        f"[bg][p]overlay=x=760+200*sin(t*1.5):y=440:shortest=1,"
        f"drawbox=x=1800:y=0:w=120:h=1080:color=white:t=fill[v]"
    )
    _ff("-filter_complex", graph, "-map", "[v]", *X264, *BT709, str(out))
    return {"fps": "25/1", "nb_frames": dur * fps, "key_rgb": GREEN, "wall_box": [1800, 0, 1920, 1080],
            "person_size": [pw, ph], "person_y": 440}


def av_offset(out: Path) -> dict[str, Any]:
    """2. MOV + PCM, video audiodan 52 ms keyin boshlanadi. Canon kabi yuvj420p (to'liq diapazon)."""
    dur, fps = 12, 25
    # -itsoffset lavfi'da 1/25 ga yaxlitlanadi (40 ms chiqadi). Shuning uchun 1/25000 time base'da butun son
    # siljish (1300 = 52 ms), -copyts (boshlanish 0 ga tortilmasin), encoder/trek/movie timescale ham 25000.
    _ff("-f", "lavfi", "-i", f"testsrc2=s={HD}:r={fps}:d={dur},settb=1/25000,setpts=PTS+1300",
        "-f", "lavfi", "-i", f"sine=f=440:r=48000:d={dur}",
        "-map", "0:v", "-map", "1:a", "-fps_mode", "passthrough", "-copyts",
        "-c:v", "libx264", "-preset", "veryfast", "-crf", "18", "-bf", "0",
        "-pix_fmt", "yuvj420p", "-color_range", "pc", *BT709,
        "-enc_time_base:v", "1/25000", "-video_track_timescale", "25000", "-movie_timescale", "25000",
        "-c:a", "pcm_s16le", str(out))
    return {"fps": "25/1", "nb_frames": dur * fps, "video_minus_audio_start_ms": 52, "pix_fmt": "yuvj420p",
            "color_range": "pc", "audio_codec": "pcm_s16le"}


def ntsc(out: Path) -> dict[str, Any]:
    """3. 29.97 fps MP4 + AAC (Fraction va AAC priming)."""
    n = 360
    _ff("-f", "lavfi", "-i", f"testsrc2=s={HD}:r=30000/1001",
        "-f", "lavfi", "-i", "sine=f=440:r=48000",
        "-frames:v", str(n), "-shortest", *X264, *BT709, "-c:a", "aac", "-b:a", "128k", str(out))
    return {"fps": "30000/1001", "nb_frames": n, "audio_codec": "aac"}


def vfr(out: Path) -> dict[str, Any]:
    """4. VFR: 25 fps manbadan har 50-kadrdan keyin bitta kadr tushirilgan (300 -> 294), pts o'zgarmagan."""
    src, fps = 300, 25
    _ff("-f", "lavfi", "-i", f"testsrc2=s={HD}:r={fps}", "-f", "lavfi", "-i", "sine=f=440:r=48000:d=12",
        "-map", "0:v", "-map", "1:a", "-vf", "select=not(eq(mod(n\\,50)\\,49))", "-frames:v", str(src - src // 50),
        "-fps_mode", "passthrough", *X264, "-c:a", "aac", str(out))
    # oxirgi kadr (299) ham tushgan: video 11.96 s -> CFR'da 299 kadr
    return {"fps": "25/1", "nb_frames": src - src // 50, "cfr_frames": src - 1, "dropped_every": 50}


def beep(out: Path) -> dict[str, Any]:
    """5. Ma'lum vaqtda 1 kHz, 300 ms, Hann oynali beep (TZ 9.6), past shovqin fonida. MOV + PCM."""
    dur, at_ms = 10, 3217
    tone = "aevalsrc=0.5*sin(2*PI*1000*t)*0.5*(1-cos(2*PI*t/0.3)):s=48000:d=0.3"
    graph = (f"{tone},adelay={at_ms}:all=1[b];anoisesrc=a=0.01:r=48000:d={dur}:seed=1[n];"
             f"[n][b]amix=inputs=2:duration=first:normalize=0[a]")
    _ff("-f", "lavfi", "-i", f"color=c=black:s=640x360:r=25:d={dur}",
        "-filter_complex", graph, "-map", "0:v", "-map", "[a]", *X264, "-c:a", "pcm_s16le", str(out))
    return {"beep_start_ms": at_ms, "beep_freq": 1000, "beep_ms": 300}


# Kadr raqami piksel kodi: yuqori chap burchakda 16 ta 64x64 blok, i-blok = N ning i-biti (oq=1, qora=0)
TIMER_BITS, TIMER_BLOCK = 16, 64


def timer(out: Path) -> dict[str, Any]:
    """6. Har kadrda o'z raqami (piksel kodi) — kadr aniqligi testlari uchun, OCR'siz."""
    n, fps = 250, 25
    w, b = TIMER_BITS * TIMER_BLOCK, TIMER_BLOCK
    code = (f"color=c=black:s={w}x{b}:r={fps},"
            f"geq=lum=16+219*mod(floor(N/pow(2\\,floor(X/{b})))\\,2):cb=128:cr=128,"
            f"pad=1920:1080:0:0:gray")
    _ff("-f", "lavfi", "-i", code, "-frames:v", str(n), *X264, *BT709, str(out))
    return {"fps": "25/1", "nb_frames": n, "bits": TIMER_BITS, "block": TIMER_BLOCK}


SLIDES = [
    ("Kirish", "Bugungi dars rejasi va maqsadlar"),
    ("Funksiya tushunchasi", "Funksiya har bir x ga bitta y ni mos qoʻyadi"),
    ("Aniqlanish sohasi", "Funksiya aniqlangan barcha x qiymatlar toʻplami"),
    ("Qiymatlar sohasi", "Funksiya qabul qiladigan barcha y qiymatlar"),
    ("Chiziqli funksiya", "y = kx + b, k burchak koeffitsiyenti"),
    ("Grafik chizish", "Ikki nuqta orqali toʻgʻri chiziq oʻtkaziladi"),
    ("Kvadrat funksiya", "y = ax² + bx + c, grafigi parabola"),
    ("Parabola uchi", "Uchning abssissasi x = −b / 2a"),
    ("Misol", "y = x² − 4x + 3 funksiyasini tekshiring"),
    ("Mashqlar", "Darslikdagi 5–9 mashqlarni bajaring"),
    ("Xulosa", "Funksiya turlari va ularning grafiklari"),
    ("Uyga vazifa", "12-mashq, 3-bet"),
]
ANIMATED_SLIDE = 6  # 1-indeksli


def _add_entrance_animation(slide: Any, shape_id: int) -> None:
    """python-pptx animatsiya API'siga ega emas: oddiy "Appear" kirish effektini XML bilan qo'shamiz."""
    from lxml import etree

    xml = f"""<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">
<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr>
<p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
<p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>
</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""
    slide._element.append(etree.fromstring(xml.replace("\n", "")))


def pptx(out: Path) -> dict[str, Any]:
    """7. 12 slaydli PPTX: ma'lum sarlavha/matn, eslatmalar, bitta animatsiyali slayd."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    layout = prs.slide_layouts[1]  # Title and Content
    expected = []
    for i, (title, body) in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
        notes = f"{i}-slayd eslatmasi: {title.lower()} haqida gapiriladi"
        s.notes_slide.notes_text_frame.text = notes
        if i == ANIMATED_SLIDE:
            _add_entrance_animation(s, s.placeholders[1].shape_id)
        expected.append({"index": i, "title": title, "body": body, "notes": notes, "animated": i == ANIMATED_SLIDE})
    prs.save(str(out))
    return {"slides": expected}


CLIPS: dict[str, Callable[[Path], dict[str, Any]]] = {
    "green_key.mp4": green_key,
    "av_offset.mov": av_offset,
    "ntsc_2997.mp4": ntsc,
    "vfr.mp4": vfr,
    "beep.mov": beep,
    "timer.mp4": timer,
    "lesson.pptx": pptx,
}


def generate(out_dir: Path = DEFAULT_OUT) -> dict[str, Any]:
    """To'plamni yaratadi (versiya mos bo'lsa mavjudini qaytaradi). manifest: {nom: {path, ...kutilgan}}."""
    out_dir.mkdir(parents=True, exist_ok=True)
    mf = out_dir / "manifest.json"
    if mf.exists():
        m: dict[str, Any] = json.loads(mf.read_text(encoding="utf-8"))
        if m.get("version") == VERSION and all((out_dir / n).exists() for n in CLIPS):
            return m
    m = {"version": VERSION, "files": {}}
    for name, fn in CLIPS.items():
        m["files"][name] = {"path": str(out_dir / name), **fn(out_dir / name)}
    tmp = mf.with_suffix(".tmp")
    tmp.write_text(json.dumps(m, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(mf)
    return m


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    for k, v in generate(out)["files"].items():
        sys.stdout.write(f"{k}: {v['path']}\n")
