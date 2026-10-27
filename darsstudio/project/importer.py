"""Loyiha yaratish va import (FR-01). Manba fayllar ko'chirilmaydi (havola + hash), so'ralsa nusxalanadi."""

import logging
import shutil
import uuid
from collections.abc import Callable
from pathlib import Path
from typing import Any

from darsstudio.media import ffmpeg, normalize, probe, proxy
from darsstudio.media.probe import MediaError
from darsstudio.project.model import MediaAsset, Project, Slide, VideoPart
from darsstudio.project.store import ProjectStore

log = logging.getLogger(__name__)

VIDEO_EXT = (".mov", ".mp4", ".mkv")
SLIDE_EXT = (".pptx", ".pdf")


def _same_format(a: probe.ProbeResult, b: probe.ProbeResult) -> bool:
    va, vb = a.video, b.video
    same_v = (va.codec, va.width, va.height, va.fps, va.pix_fmt) == (vb.codec, vb.width, vb.height, vb.fps, vb.pix_fmt)
    au = (a.audio.codec, a.audio.sample_rate, a.audio.channels) if a.audio else None
    bu = (b.audio.codec, b.audio.sample_rate, b.audio.channels) if b.audio else None
    return same_v and au == bu


def join_parts(paths: list[Path], out: Path) -> tuple[probe.ProbeResult, list[VideoPart]]:
    """Qismlarni qayta kodlashsiz ketma-ket ulaydi (concat demuxer, -c copy): sifat yo'qolmaydi, kadr indekslari
    qismlar bo'yicha uzluksiz. Format (kodek, o'lcham, fps, audio) bir xil bo'lishi shart."""
    probes = [probe.probe(p) for p in paths]
    for p, pr in zip(paths, probes, strict=True):
        probe.validate(pr)
        if normalize.needs_normalize(pr):
            raise MediaError(f"{p.name}: o'zgaruvchan fps / boshqa audio chastotali qismlarni ulash hozircha "
                             "qo'llanmaydi — qismlarni bitta formatda yozing.")
        if not _same_format(probes[0], pr):
            raise MediaError(f"{p.name} formati {paths[0].name} dan farq qiladi (o'lcham/fps/kodek): "
                             "qismlar bir xil kamera sozlamasida yozilishi kerak.")
    out.parent.mkdir(parents=True, exist_ok=True)
    lst = out.with_suffix(".txt")
    lst.write_text("".join(f"file '{p.resolve().as_posix()}'\n" for p in paths), encoding="utf-8")
    tmp = out.with_name(out.stem + ".tmp" + out.suffix)
    ffmpeg.run(["-v", "error", "-nostdin", "-y", "-f", "concat", "-safe", "0", "-i", str(lst), "-map", "0:v:0",
                *(["-map", "0:a:0"] if probes[0].audio else []), "-c", "copy", str(tmp)])
    tmp.replace(out)
    joined = probe.probe(out)
    parts, start = [], 0
    for p, pr in zip(paths, probes, strict=True):
        n = probe.frame_count(pr)
        parts.append(VideoPart(path=str(p.resolve()), start=start, frames=n))
        start += n
    if abs(probe.frame_count(joined) - start) > 1:
        raise MediaError(f"Qismlarni ulashda kadrlar soni mos kelmadi ({probe.frame_count(joined)} != {start}).")
    return joined, parts


def append_part(store: ProjectStore, video: Path) -> Project:
    """Mavjud loyihaga davomini qo'shish. 1-qism kadr indekslari o'zgarmaydi (inson qarorlari — tasdiqlangan
    kesishlar, Eye Contact o'chirilgan joylar — saqlanadi); tahlil (audio, ASR, sinxron) qayta bajarilishi kerak."""
    from darsstudio.project.commands import project_lock
    p = store.load()
    v = p.asset("video")
    assert v is not None
    old = [Path(x.path) for x in p.parts] if p.parts else [Path(v.path)]
    joined, parts = join_parts([*old, Path(video).resolve()], store.cache("media") / "joined.mov")  # uzoq — qulfsiz
    total = probe.frame_count(joined)
    with project_lock(store.root):
        p = store.load()
        for i, o in enumerate(p.parts):   # avvalgi qismlarning key/joylashuvi saqlanadi
            parts[i] = parts[i].model_copy(update={"key": o.key, "placement": o.placement})
        keep_cuts = [c for c in p.cuts if c.status != "taklif" and c.reason not in ("bosh", "oxir", "qism")]
        assets = [a.model_copy(update={"cfr_path": str(joined.path)}) if a.kind == "video" else a for a in p.assets]
        upd: dict[str, Any] = {"parts": parts, "total_frames": total, "assets": assets, "cuts": keep_cuts}
        for name in ("slide_spans", "layout_spans"):   # oxirgi oraliq yangi oxirgacha (sinxron qayta hisoblaydi)
            spans = getattr(p, name)
            upd[name] = [*spans[:-1], spans[-1].model_copy(update={"end": total})] if spans else []
        store.save(Project.model_validate(p.model_copy(update=upd).model_dump()))
    return store.load()


def read_slide_titles(path: Path) -> list[Slide]:
    """Slaydlar soni va sarlavhalar. To'liq ekstraksiya (matn, eslatmalar, render) — slides/ modulida."""
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation
        try:
            prs = Presentation(str(path))
        except Exception as e:
            raise MediaError(f"{path.name} ni ochib bo'lmadi: PPTX buzilgan yoki parol bilan himoyalangan.") from e
        out = []
        for i, s in enumerate(prs.slides, 1):
            t = s.shapes.title
            out.append(Slide(index=i, title=(t.text if t is not None and t.has_text_frame else "").strip()))
        return out
    if path.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(path))
        return [Slide(index=i + 1) for i in range(len(doc))]
    raise MediaError(f"{path.name}: slaydlar PPTX yoki PDF bo'lishi kerak.")


def import_project(video: Path | list[Path], slides: Path, template_id: str, workdir: Path, *, name: str | None = None,
                   meta: dict[str, str] | None = None, slide_log: Path | None = None, copy: bool = False,
                   make_proxy: bool = True, on_step: Callable[[str], None] = lambda s: None) -> ProjectStore:
    videos = [Path(x).resolve() for x in (video if isinstance(video, list) else [video])]
    video, slides = videos[0], Path(slides).resolve()
    name = name or video.stem
    store = ProjectStore(Path(workdir) / name)
    if store.exists():
        raise MediaError(f"'{name}' loyihasi allaqachon bor: {store.root}. Boshqa nom bering yoki eskisini o'chiring.")

    on_step("probe")
    p = probe.probe(video)
    probe.validate(p)
    slide_list = read_slide_titles(slides)
    if not slide_list:
        raise MediaError(f"{slides.name} da slayd yo'q.")

    store.create_dirs()
    try:
        if copy:
            src_dir = store.root / "source"
            src_dir.mkdir(exist_ok=True)
            for f in (video, slides, *([slide_log] if slide_log else [])):
                shutil.copy2(f, src_dir / f.name)
            video, slides = src_dir / video.name, src_dir / slides.name
            slide_log = src_dir / slide_log.name if slide_log else None

        work = p
        cfr_path = None
        parts: list[VideoPart] = []
        if len(videos) > 1:
            on_step("join")
            work, parts = join_parts(videos, store.cache("media") / "joined.mov")
            cfr_path = work.path
        elif normalize.needs_normalize(p):
            on_step("normalize")
            work = normalize.normalize(p, store.cache("media") / "cfr.mkv")
            cfr_path = work.path
            log.info("normallashtirildi: %s -> %s (vfr=%s)", video, cfr_path, p.video.vfr)

        assets = [
            MediaAsset(id=probe.file_id(video), path=str(video), kind="video", probe=p.to_dict(), cfr_path=cfr_path),
            MediaAsset(id=probe.file_id(slides), path=str(slides), kind=slides.suffix.lower().lstrip(".")),
        ]
        if slide_log:
            assets.append(MediaAsset(id=probe.file_id(slide_log), path=str(Path(slide_log).resolve()),
                                     kind="slide_log"))

        project = Project(id=uuid.uuid4().hex[:12], name=name, meta=meta or {}, assets=assets,
                          fps=str(work.video.fps), total_frames=probe.frame_count(work), template_id=template_id,
                          slides=slide_list, parts=parts)
        store.save(project)
        if make_proxy:
            on_step("proxy")
            proxy.make_proxy(work, store.cache("media") / "proxy_540p.mp4")
    except BaseException:
        shutil.rmtree(store.root, ignore_errors=True)  # yarim loyiha qolmasin: qayta import mumkin bo'lsin
        raise
    return store


def find_pairs(folder: Path) -> list[tuple[Path, Path]]:
    """Papkadagi juftliklar: dars05.mov + dars05.pptx (bir xil nom, katta-kichik harf farqsiz)."""
    files = [f for f in Path(folder).iterdir() if f.is_file()]
    slides = {f.stem.lower(): f for f in files if f.suffix.lower() in SLIDE_EXT}
    return sorted((v, slides[v.stem.lower()]) for v in files
                  if v.suffix.lower() in VIDEO_EXT and v.stem.lower() in slides)
