"""PPTX matni va ogohlantirishlar (FR-03, TZ 4.2): sarlavha, asosiy matn, jadval kataklari, ma'ruzachi eslatmalari.
Ogohlantirishlar: animatsiya (yakuniy holat ko'rsatiladi), video/audio (statik kadr), tizimda yo'q shriftlar."""

import functools
import os
import subprocess
import sys
from collections.abc import Iterator
from typing import Any

from darsstudio.project.model import Slide

P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
WARN_ANIM = "animatsiya: slaydning yakuniy holati ko'rsatiladi"
WARN_ANIM_STEPS = "animatsiya: {n} qadam (javoblar o'qituvchi aytganda chiqadi)"
WARN_MEDIA = "video/audio: statik kadr ko'rsatiladi"


def _shapes(shapes: Any) -> Iterator[Any]:
    for sh in shapes:
        if sh.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _shapes(sh.shapes)
        else:
            yield sh


def _text(sh: Any) -> str:
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        return "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
    if getattr(sh, "has_table", False) and sh.has_table:
        return "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in sh.table.rows)
    return ""


def anim_clicks(slide_el: Any) -> tuple[list[int], list[tuple[list[int], list[int]]]]:
    """Asosiy animatsiya ketma-ketligi: (boshidan avtomatik chiqadigan shakllar, har bosish uchun
    (paydo bo'ladigan, yo'qoladigan) shakl id'lari). Ta'kidlash (emph) e'tiborga olinmaydi."""
    timing = slide_el.find(f"{P_NS}timing")
    if timing is None:
        return [], []
    seq = next((c for c in timing.iter(f"{P_NS}cTn") if c.get("nodeType") == "mainSeq"), None)
    kids = seq.find(f"{P_NS}childTnLst") if seq is not None else None
    if kids is None:
        return [], []
    auto: list[int] = []
    clicks: list[tuple[list[int], list[int]]] = []
    for par in kids:
        show: list[int] = []
        hide: list[int] = []
        click = False
        for c in par.iter(f"{P_NS}cTn"):
            click |= c.get("nodeType") == "clickEffect"
            cls = c.get("presetClass")
            tgt = c.find(f".//{P_NS}spTgt")
            if cls in ("entr", "exit") and tgt is not None and tgt.get("spid", "").isdigit():
                (show if cls == "entr" else hide).append(int(tgt.get("spid")))
        if not show and not hide:
            continue
        if click or clicks:
            clicks.append((show, hide))
        else:
            auto += show
    return auto, clicks


def _ends(sh: Any) -> list[tuple[float, float]]:
    """Chiziq/strelka uchlari (flip hisobga olingan); boshqa shakl uchun markazi."""
    x, y, w, h = (sh.left or 0), (sh.top or 0), (sh.width or 0), (sh.height or 0)
    xf = sh._element.find(f".//{A_NS}xfrm")
    fh = xf is not None and xf.get("flipH") == "1"
    fv = xf is not None and xf.get("flipV") == "1"
    geom = sh._element.find(f".//{A_NS}prstGeom")
    line = geom is not None and ("line" in geom.get("prst", "").lower() or "connector" in geom.get("prst", "").lower())
    if line or sh.shape_type in (9, 21) or sh._element.tag.endswith("cxnSp"):
        a, b = (x + w, x) if fh else (x, x + w)
        c, d = (y + h, y) if fv else (y, y + h)
        return [(a, c), (b, d)]
    return [(x + w / 2, y + h / 2)]


def anim_texts(slide: Any, clicks: list[tuple[list[int], list[int]]]) -> list[str]:
    """Har bosishda chiqadigan matn. Matnsiz shakl (moslashtirish chizig'i) — uchiga eng yaqin matnli shakl
    (ko'p so'zlisi: "1. loyalty" -> "d) the tendency to always buy..." chizig'ida ta'rif)."""
    shapes = list(_shapes(slide.shapes))
    by_id = {sh.shape_id: sh for sh in shapes}
    animated = {i for show, _ in clicks for i in show}
    static = [(sh, txt) for sh in shapes if sh.shape_id not in animated and (txt := _text(sh).strip())]

    def near(pt: tuple[float, float]) -> str:
        def dist(sh: Any) -> float:
            x0, y0 = sh.left or 0, sh.top or 0
            dx = max(x0 - pt[0], 0, pt[0] - x0 - (sh.width or 0))
            dy = max(y0 - pt[1], 0, pt[1] - y0 - (sh.height or 0))
            return float(dx * dx + dy * dy)
        return min(static, key=lambda st: dist(st[0]))[1] if static else ""

    out = []
    for show, _ in clicks:
        shs = [by_id[i] for i in show if i in by_id]
        parts = [txt for sh in shs if (txt := _text(sh).strip())]
        if not parts and static:   # faqat chiziqlar (javob qutisi foni emas)
            for sh in shs:
                ends = _ends(sh)
                txt = max((near(pt) for pt in ends), key=lambda s: len(s.split())) if len(ends) == 2 else ""
                if txt and txt not in parts:
                    parts.append(txt)
        out.append(" ".join(parts))
    return out


def _fonts(slide: Any) -> set[str]:
    out = set()
    for el in slide._element.iter(f"{A_NS}latin", f"{A_NS}cs", f"{A_NS}ea"):
        tf = el.get("typeface")
        if tf and not tf.startswith("+"):  # +mj-lt/+mn-lt — mavzu shriftlari, alohida tekshiriladi
            out.add(tf)
    return out


def _theme_fonts(prs: Any) -> set[str]:
    out = set()
    for master in prs.slide_masters:
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                from lxml import etree
                root = etree.fromstring(rel.target_part.blob)
                for tag in ("majorFont", "minorFont"):
                    for el in root.iter(f"{A_NS}{tag}"):
                        latin = el.find(f"{A_NS}latin")
                        if latin is not None and latin.get("typeface"):
                            out.add(latin.get("typeface"))
    return out


@functools.cache
def installed_fonts() -> frozenset[str]:
    """Tizimdagi shrift oilalari (kichik harfda). Windows: registry; boshqasi: fc-list."""
    names: set[str] = set()
    if sys.platform == "win32":
        import winreg
        for hive, key in ((winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
                          (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")):
            try:
                with winreg.OpenKey(hive, key) as k:
                    i = 0
                    while True:
                        try:
                            n = winreg.EnumValue(k, i)[0]
                        except OSError:
                            break
                        i += 1
                        n = n.split("(")[0].strip()          # "Arial Bold (TrueType)" -> "Arial Bold"
                        for part in n.split("&"):
                            names.add(part.strip().lower())
            except OSError:
                pass
    else:
        try:
            r = subprocess.run(["fc-list", ":", "family"], capture_output=True, encoding="utf-8", errors="replace")
            for ln in r.stdout.splitlines():
                names.update(x.strip().lower() for x in ln.split(","))
        except FileNotFoundError:
            pass
    return frozenset(names)


def font_missing(name: str) -> bool:
    fonts = installed_fonts()
    if not fonts:
        return False  # tekshirib bo'lmadi — ogohlantirmaymiz
    n = name.lower()
    return not any(f == n or f.startswith(n + " ") for f in fonts)


def extract_pptx(path: str | os.PathLike[str]) -> tuple[list[Slide], list[str]]:
    """(slaydlar, umumiy ogohlantirishlar). png_4k bo'sh qoladi — render modullari to'ldiradi."""
    from pptx import Presentation

    prs = Presentation(str(path))
    general = []
    if prs.slide_width and prs.slide_height and abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        general.append(f"slayd nisbati {prs.slide_width / prs.slide_height:.2f} (16:9 emas): chetlari qora (pillarbox)")
    theme = _theme_fonts(prs)
    out = []
    for i, s in enumerate(prs.slides, 1):
        title_shape = s.shapes.title
        title = _text(title_shape).strip() if title_shape is not None else ""
        tid = title_shape.shape_id if title_shape is not None else None  # proxy obyektlar: id bilan
        if not title:
            # sarlavha placeholder'i yo'q (dizayn shablonlarida ko'p): eng yuqoridagi qisqa (<= 12 so'z) matn
            cands = [(sh.top or 0, sh.shape_id, t) for sh in _shapes(s.shapes)
                     if (t := _text(sh).strip()) and len(t.split()) <= 12 and "\n" not in t]
            if cands:
                _, tid, title = min(cands)
        body = [t for sh in _shapes(s.shapes) if sh.shape_id != tid and (t := _text(sh).strip())]
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = s.notes_slide.notes_text_frame.text.strip()
        warns = []
        timing = s._element.find(f"{P_NS}timing")
        _, clicks = anim_clicks(s._element)
        anim = anim_texts(s, clicks)
        if clicks:
            warns.append(WARN_ANIM_STEPS.format(n=len(clicks)))
        elif timing is not None and any(c.get("presetClass") for c in timing.iter(f"{P_NS}cTn")):
            warns.append(WARN_ANIM)
        if any(sh.shape_type == 16 for sh in _shapes(s.shapes)):  # MSO_SHAPE_TYPE.MEDIA
            warns.append(WARN_MEDIA)
        missing = sorted(f for f in _fonts(s) | theme if font_missing(f))
        if missing:
            warns.append("tizimda yo'q shrift: " + ", ".join(missing))
        out.append(Slide(index=i, title=title, body="\n".join(body), notes=notes, warnings=warns, anim=anim))
    return out, general
