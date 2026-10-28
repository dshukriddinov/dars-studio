"""Slaydlarni 4K PNG ga render qilish (FR-03, TZ 4.2).

PPTX: PowerPoint COM (Slide.Export) -> bo'lmasa LibreOffice headless -> PDF -> pypdfium2. PDF: pypdfium2.
PNG lar slaydning o'z nisbatida, balandligi 2160 (16:9 da 3840x2160); 16:9 ga pillarbox — warp/kompozitda.
PyMuPDF ishlatilmaydi (AGPL, TZ 10.2). PowerPoint bir vaqtda bitta jarayonda (rejalashtiruvchida `pptx` = 1).
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from darsstudio.media.probe import MediaError, file_id

HEIGHT = 2160


def powerpoint_available() -> bool:
    if sys.platform != "win32":
        return False
    import winreg
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application"))
        return True
    except OSError:
        return False


def soffice_path() -> str | None:
    for p in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if p and Path(p).is_file():
            return p
    return None


def _com_shape(shapes: Any, sid: int) -> Any:
    """Slayd shakli Id bo'yicha (guruh ichida ham)."""
    for i in range(1, shapes.Count + 1):
        sh = shapes(i)
        if sh.Id == sid:
            return sh
        if sh.Type == 6:  # msoGroup
            for j in range(1, sh.GroupItems.Count + 1):
                if sh.GroupItems(j).Id == sid:
                    return sh.GroupItems(j)
    return None


def anim_states(src: Path) -> dict[int, list[set[int]]]:
    """Animatsiyali slaydlar: {slayd: [k bosishdan keyin yashirin shakllar, k = 0..n-1]} (n — bosishlar soni)."""
    from pptx import Presentation

    from darsstudio.slides.extract import anim_clicks
    out: dict[int, list[set[int]]] = {}
    for i, s in enumerate(Presentation(str(src)).slides, 1):
        _, clicks = anim_clicks(s._element)
        states = []
        for k in range(len(clicks)):
            later = {x for show, _ in clicks[k:] for x in show}      # hali chiqmagan
            gone = {x for _, hide in clicks[:k] for x in hide}       # yo'qolgan
            states.append(later | gone)
        if states:
            out[i] = states
    return out


def _export_states(sl: Any, states: list[set[int]], stem: Path, w: int) -> list[Path]:
    """Har holat: hali chiqmagan shakllar yashirilib eksport, keyin qaytariladi. COM havolalari shu funksiyada
    qoladi (PowerPoint yopilgandan keyin bo'shatilsa RPC_E_DISCONNECTED)."""
    out = []
    for k, hidden in enumerate(states):
        shs = [s for s in (_com_shape(sl.Shapes, x) for x in hidden) if s is not None]
        for s in shs:
            s.Visible = 0
        q = stem.with_name(f"{stem.name}_{k}.png")
        sl.Export(str(q.resolve()), "PNG", w, HEIGHT)
        for s in shs:
            s.Visible = -1
        out.append(q)
    return out


def render_pptx_com(src: Path, out_dir: Path) -> tuple[list[Path], dict[int, list[Path]]]:
    """(yakuniy holatlar, {slayd: animatsiya qadamlari 0..n-1}). Qadam — shakllar xotirada yashirilib eksport."""
    import pythoncom
    import win32com.client

    states = anim_states(src)
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    pres = None
    try:
        # WithWindow=False: foydalanuvchi oynasi ochilmaydi; ReadOnly: manba o'zgarmaydi
        pres = app.Presentations.Open(str(src.resolve()), True, False, False)
        w = round(HEIGHT * pres.PageSetup.SlideWidth / pres.PageSetup.SlideHeight)
        out = []
        steps: dict[int, list[Path]] = {}
        for i in range(1, pres.Slides.Count + 1):
            p = out_dir / f"{i:03d}.png"
            pres.Slides(i).Export(str(p.resolve()), "PNG", w, HEIGHT)
            out.append(p)
            if i in states:
                steps[i] = _export_states(pres.Slides(i), states[i], out_dir / f"{i:03d}", w)
        pres.Saved = True   # o'zgarishlar faqat xotirada; manba saqlanmaydi
        return out, steps
    finally:
        if pres is not None:
            pres.Close()
        if app.Presentations.Count == 0:  # foydalanuvchi PowerPoint'da boshqa fayl ochgan bo'lsa yopmaymiz
            app.Quit()
        pythoncom.CoUninitialize()


def render_pdf(src: Path, out_dir: Path) -> list[Path]:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(src))
    out = []
    try:
        for i in range(len(doc)):
            page = doc[i]
            _, h = page.get_size()
            img = page.render(scale=HEIGHT / h).to_pil().convert("RGB")
            p = out_dir / f"{i + 1:03d}.png"
            img.save(p)
            out.append(p)
    finally:
        doc.close()
    return out


def pptx_to_pdf_libreoffice(src: Path, tmp: Path) -> Path:
    exe = soffice_path()
    if exe is None:
        raise MediaError("PPTX ni ochish uchun PowerPoint yoki LibreOffice kerak. Ulardan birini o'rnating "
                         "yoki slaydlarni PDF qilib bering.")
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)], check=True,
                   capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    return tmp / (src.stem + ".pdf")


def render_slides(src: Path, out_dir: Path) -> tuple[list[Path], str]:
    """(PNG ro'yxati, ishlatilgan usul). Manba hash'i mos bo'lsa keshdan qaytaradi."""
    src = Path(src)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = out_dir / "stamp.json"
    sid = file_id(src)
    if stamp.is_file():
        st = json.loads(stamp.read_text(encoding="utf-8"))
        files = [out_dir / f for f in st.get("files", [])]
        fresh = "steps" in st or st.get("method") != "powerpoint"   # eski kesh: animatsiya qadamlarisiz
        if st.get("source_id") == sid and files and all(f.is_file() for f in files) and fresh:
            return files, st["method"]
    for old in out_dir.glob("*.png"):
        old.unlink()
    ext = src.suffix.lower()
    steps: dict[int, list[Path]] = {}
    if ext == ".pdf":
        files, method = render_pdf(src, out_dir), "pdf"
    elif ext == ".pptx" and powerpoint_available():
        (files, steps), method = render_pptx_com(src, out_dir), "powerpoint"
    elif ext == ".pptx":
        with tempfile.TemporaryDirectory() as t:
            files, method = render_pdf(pptx_to_pdf_libreoffice(src, Path(t)), out_dir), "libreoffice"
    else:
        raise MediaError(f"{src.name}: slaydlar PPTX yoki PDF bo'lishi kerak.")
    stamp.write_text(json.dumps({"source_id": sid, "method": method, "files": [f.name for f in files],
                                 "steps": {str(i): [f.name for f in fs] for i, fs in steps.items()}}),
                     encoding="utf-8")
    return files, method


def load_steps(out_dir: Path) -> dict[int, list[Path]]:
    """render_slides natijasidagi animatsiya qadamlari (PowerPoint bilan render qilingan bo'lsa)."""
    stamp = out_dir / "stamp.json"
    if not stamp.is_file():
        return {}
    st = json.loads(stamp.read_text(encoding="utf-8"))
    return {int(i): [out_dir / f for f in fs] for i, fs in st.get("steps", {}).items()}
